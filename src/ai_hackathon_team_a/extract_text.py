"""文字抽出（設計書 §4 の3）。

拡張子の許可リストと ``.xlsm``・古い形式の拒否は呼び出し側（API ルーター）の役目。ここは
``.xlsx``・``.docx``・``.pptx``・``.pdf``・その他（UTF-8）の抽出だけを担当する。
"""

import csv
import io
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

import docx
import openpyxl
import pptx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader


class ExtractionError(Exception):
    """ファイルの中身を読み取れなかったことを表す（呼び出し側で400相当に変換する）。"""


@dataclass(frozen=True)
class ExtractedContent:
    """文字抽出の結果。

    ``rows`` は ``.xlsx``（シートの行）と ``.pptx``（スライドとノート）のときに設定する。
    (locator, 文字) の並びで、1つが1区切りになる（伏せ字も単位ごとに行う）。
    ``blocks`` は ``.docx`` のときに設定する。(段落番号, 1行の文字) の並びで、空の文字は
    空行（区切りの境目）を表す。区切りはテキストと同じ規則（空行か20行ごと）。
    どちらも ``None`` なら ``text`` 全体をテキストの規則で区切る。
    """

    text: str
    rows: list[tuple[str, str]] | None = None
    blocks: list[tuple[int, str]] | None = None


def extract_text(filename: str, data: bytes) -> ExtractedContent:
    """ファイルの中身から文字を取り出す。読めない場合は :class:`ExtractionError`。"""

    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        return _extract_xlsx(data)
    if suffix == ".docx":
        return _extract_docx(data)
    if suffix == ".pptx":
        return _extract_pptx(data)
    if suffix == ".pdf":
        return _extract_pdf(data)
    return _extract_plain_text(data)


def _extract_xlsx(data: bytes) -> ExtractedContent:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ExtractionError("Excel ファイルを読み取れませんでした。") from exc

    rows: list[tuple[str, str]] = []
    text_lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            text_lines.append(f"# {sheet.title}")
            for row_no, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                if all(v is None or str(v).strip() == "" for v in values):
                    continue  # 空の行は区切りにしない（「,,」だけの区切りを作らない）
                csv_line = _row_to_csv(values)
                rows.append((f"{sheet.title}:{row_no}", csv_line))
                text_lines.append(csv_line)
    finally:
        workbook.close()

    return ExtractedContent(text="\n".join(text_lines), rows=rows)


def _extract_docx(data: bytes) -> ExtractedContent:
    """段落と表を文書の順に抜く。見出しは「## 見出し」、表は1行を「セル1 | セル2 | …」に。

    本文の段落と表の行に、文書の頭から1、2、…と番号を振る（locator の段落番号）。
    ヘッダー・フッター・コメントは対象外（本文 ``document.element.body`` だけを読む）。
    """

    try:
        document = docx.Document(io.BytesIO(data))
        blocks: list[tuple[int, str]] = []
        number = 0
        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                number += 1
                paragraph = DocxParagraph(child, document)
                text = paragraph.text.strip()
                style = (paragraph.style.name if paragraph.style is not None else "") or ""
                if text and (
                    style.startswith("Heading") or style.startswith("見出し") or style == "Title"
                ):
                    text = f"## {text}"
                blocks.append((number, text))
            elif tag == "tbl":
                for row in DocxTable(child, document).rows:
                    number += 1
                    cells = [_one_line(cell.text) for cell in row.cells]
                    line = " | ".join(cells)
                    blocks.append((number, line if any(cells) else ""))
                blocks.append((number, ""))  # 表のあとは段落の境目にする
    except Exception as exc:
        raise ExtractionError("Word ファイルを読み取れませんでした。") from exc
    return ExtractedContent(text="\n".join(t for _, t in blocks), blocks=blocks)


def _extract_pptx(data: bytes) -> ExtractedContent:
    """スライドごとに、タイトル・本文・表を1つに、発表者ノートは別の単位にする。

    グループ化された図形の中の文字も拾う。画像の中の文字は対象外。
    """

    try:
        presentation = pptx.Presentation(io.BytesIO(data))
        rows: list[tuple[str, str]] = []
        for number, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title
            lines: list[str] = []
            if (
                title_shape is not None
                and title_shape.has_text_frame
                and title_shape.text_frame.text.strip()
            ):
                lines.append(f"## {_one_line(title_shape.text_frame.text)}")
            for shape in slide.shapes:
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                lines.extend(_shape_lines(shape))
            body = "\n".join(line for line in lines if line.strip())
            if body:
                rows.append((f"スライド{number}", body))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                notes_text = notes.text.strip() if notes is not None else ""
                if notes_text:
                    rows.append((f"スライド{number}のノート", notes_text))
    except Exception as exc:
        raise ExtractionError("PowerPoint ファイルを読み取れませんでした。") from exc
    return ExtractedContent(text="\n\n".join(t for _, t in rows), rows=rows)


def _shape_lines(shape) -> list[str]:
    """図形1つから文字の行を取り出す（グループは中まで、表は1行＝「セル | セル」）。"""

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [line for inner in shape.shapes for line in _shape_lines(inner)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return [" | ".join(_one_line(cell.text) for cell in row.cells) for row in shape.table.rows]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
    return []


def _one_line(text: str) -> str:
    return " ".join(text.split())


def _extract_pdf(data: bytes) -> ExtractedContent:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ExtractionError("PDF を読み取れませんでした。") from exc
    return ExtractedContent(text="\n\n".join(pages))


def _extract_plain_text(data: bytes) -> ExtractedContent:
    try:
        return ExtractedContent(text=data.decode("utf-8"))
    except UnicodeDecodeError as exc:
        raise ExtractionError("UTF-8 として読み取れませんでした。") from exc


def _row_to_csv(values: Sequence[object]) -> str:
    """1行のセルの値を CSV の1行にする（改行はスペースに寄せて1行を保つ）。"""

    cleaned = ["" if v is None else str(v).replace("\r", " ").replace("\n", " ") for v in values]
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="").writerow(cleaned)
    return buffer.getvalue()
