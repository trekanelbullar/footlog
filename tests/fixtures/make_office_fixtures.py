"""tests/fixtures の小さな Office ファイル（docx・pptx・xlsx）を作り直すスクリプト。

実在の会社名・人名は入れない。作り直すとき：``uv run python tests/fixtures/make_office_fixtures.py``
"""

from pathlib import Path

import docx
import openpyxl
import pptx
from pptx.util import Inches

HERE = Path(__file__).parent


def make_docx(path: Path) -> None:
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "ヘッダーの文字（対象外）"
    document.sections[0].footer.paragraphs[0].text = "フッターの文字（対象外）"
    document.add_heading("配信方式の検討", level=1)
    document.add_paragraph("通知メールの送信には外部サービスを使うことに決めた。")
    document.add_paragraph("理由は、自前でメールサーバーを持つと運用の手間が大きいため。")
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=3)
    for cells, values in zip(
        table.rows, [("案", "費用", "判断"), ("自前サーバー", "高い", "却下")], strict=True
    ):
        for cell, value in zip(cells.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("次は文面の確認を行う。")
    document.save(path)


def make_pptx(path: Path) -> None:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "第1回の定例"
    slide.placeholders[1].text = "デモは5分で行うことに決めた"
    slide.notes_slide.notes_text_frame.text = "時間が足りなければ質疑を短くする"

    slide2 = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide2.shapes.title.text = "残った論点"
    group = slide2.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "グループの中の文字：資料はA4一枚"
    rows = slide2.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    rows.cell(0, 0).text, rows.cell(0, 1).text = "論点", "状態"
    rows.cell(1, 0).text, rows.cell(1, 1).text = "文面", "未定"
    presentation.save(path)


def make_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "進捗"
    sheet.append(["項目", "状態", "メモ"])
    sheet.append(["デプロイ", "完了", "東京リージョン"])
    sheet.append([None, None, None])
    sheet.append(["通知", "作業中", "文面は未定"])
    other = workbook.create_sheet("決定")
    other.append(["配信は外部サービス", "理由：運用の手間"])
    workbook.save(path)


if __name__ == "__main__":
    make_docx(HERE / "sample.docx")
    make_pptx(HERE / "sample.pptx")
    make_xlsx(HERE / "sample.xlsx")
